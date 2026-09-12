"""
Multi-format plain text extraction module.
Supports PDF, DOCX, XLSX, XLS, PPTX, RTF, TXT, CSV, JSON, XML, HTML, MD, YAML, LOG.
"""
import io
import os
import re
from typing import Set
import fitz  # PyMuPDF
from PyPDF2 import PdfReader
import docx
import openpyxl
import xlrd
from pptx import Presentation
from striprtf.striprtf import rtf_to_text

TEXT_EXTRACTABLE_EXTENSIONS: Set[str] = {
    '.txt', '.csv', '.json', '.xml', '.html', '.htm', '.md',
    '.log', '.yaml', '.yml', '.ini', '.cfg', '.toml',
    '.pdf', '.docx', '.xlsx', '.xls', '.pptx', '.rtf'
}

def extract_text(file_bytes: bytes, filename: str) -> str:
    """
    Extracts plain text content from file bytes based on file extension.
    Returns empty string if the file format is non-textual or unsupported.
    """
    if not file_bytes or not filename:
        return ""

    ext = os.path.splitext(filename.lower())[1]
    if ext not in TEXT_EXTRACTABLE_EXTENSIONS:
        return ""

    try:
        if ext == '.pdf':
            # Primary: PyMuPDF (fitz)
            try:
                doc = fitz.open(stream=file_bytes, filetype="pdf")
                parts = [page.get_text("text") for page in doc if page.get_text("text").strip()]
                doc.close()
                extracted = " ".join(parts).strip()
                if extracted:
                    return extracted
            except Exception:
                pass
            # Fallback: PyPDF2
            try:
                pdf = PdfReader(io.BytesIO(file_bytes))
                return " ".join([p.extract_text() for p in pdf.pages if p.extract_text()]).strip()
            except Exception:
                return ""

        elif ext == '.docx':
            doc = docx.Document(io.BytesIO(file_bytes))
            parts = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            parts.append(cell.text)
            return " ".join(parts).strip()

        elif ext == '.xlsx':
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
            parts = []
            for ws in wb.worksheets:
                for r in ws.iter_rows(values_only=True):
                    for c in r:
                        if c is not None:
                            parts.append(str(c))
            return " ".join(parts).strip()

        elif ext == '.xls':
            wb = xlrd.open_workbook(file_contents=file_bytes)
            parts = []
            for sheet in wb.sheets():
                for r in range(sheet.nrows):
                    for c in range(sheet.ncols):
                        val = sheet.cell_value(r, c)
                        if val:
                            parts.append(str(val))
            return " ".join(parts).strip()

        elif ext == '.pptx':
            prs = Presentation(io.BytesIO(file_bytes))
            parts = []
            for slide in prs.slides:
                for s in slide.shapes:
                    if hasattr(s, 'text') and s.text.strip():
                        parts.append(s.text)
            return " ".join(parts).strip()

        elif ext == '.rtf':
            try:
                decoded = file_bytes.decode('utf-8')
            except UnicodeDecodeError:
                decoded = file_bytes.decode('latin-1', errors='replace')
            return rtf_to_text(decoded).strip()

        else:
            # Plain text, CSV, JSON, Markdown, XML, HTML, etc.
            try:
                raw = file_bytes.decode('utf-8')
            except UnicodeDecodeError:
                raw = file_bytes.decode('latin-1', errors='replace')
            if ext in ('.xml', '.html', '.htm'):
                raw = re.sub(r'<[^>]+>', ' ', raw)
            return raw.strip()

    except Exception:
        return ""
